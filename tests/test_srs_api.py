"""Các API bổ sung theo SRS bảng 8.5, 8.9, 8.10: nhập sản phẩm từ CSV, /api/inventory/*, /api/export/*."""
from tests.helpers import product_id

TEMPLATE_HEADER = "sku,ma_vach,ten,nhom_hang,hang,gia_ban,gia_von,vat,bao_hanh_thang,ton_toi_thieu,theo_serial,mo_ta\n"


def upload(client, h, text: str, dry_run=False, encoding="utf-8-sig"):
    return client.post("/api/products/import", params={"dry_run": dry_run}, headers=h,
                       files={"file": ("sp.csv", text.encode(encoding), "text/csv")})


# ---------------- Nhập sản phẩm từ CSV ----------------
def test_template_download(client, owner_h):
    r = client.get("/api/products/import-template", headers=owner_h)
    assert r.status_code == 200 and r.content.startswith("﻿".encode())  # UTF-8 có BOM
    assert r.content.decode("utf-8-sig").startswith(TEMPLATE_HEADER.strip())


def test_import_dry_run_then_create(client, owner_h):
    csv_text = TEMPLATE_HEADER + (
        "LT-ASUS-09,,Laptop Asus Vivobook 15,Laptop,Asus,\"15.990.000\",13500000,10,24,2,1,\n"
        "PK001,,Trùng SKU đã có,Phụ kiện,,100000,,,,,,\n"
        "CAP-01,,Cáp USB-C 1m,Phụ kiện,,90000,40000,,,10,0,Cáp sạc nhanh\n")
    preview = upload(client, owner_h, csv_text, dry_run=True).json()
    assert preview["to_create"] == 2 and preview["errors"] == [] and preview["skipped"][0]["sku"] == "PK001"
    assert client.get("/api/products", params={"q": "Vivobook"}, headers=owner_h).json()["total"] == 0  # chưa ghi

    r = upload(client, owner_h, csv_text)
    assert r.status_code == 200, r.text
    assert r.json()["created"] == 2
    p = client.get("/api/products", params={"q": "Vivobook"}, headers=owner_h).json()["items"][0]
    assert p["sale_price"] == 15_990_000 and p["track_serial"] and p["warranty_months"] == 24 and p["stock"] == 0
    cats = {c["name"] for c in client.get("/api/categories", headers=owner_h).json()}
    assert "Laptop" in cats  # nhóm hàng mới được tạo theo tên
    cable = client.get("/api/products", params={"q": "CAP-01"}, headers=owner_h).json()["items"][0]
    assert cable["vat_rate"] == 10  # VAT bỏ trống: lấy theo nhóm hàng (BR-04)


def test_import_is_all_or_nothing(client, owner_h):
    csv_text = TEMPLATE_HEADER + "OK-01,,Hợp lệ,,,100000,,,,,,\nBAD-01,,Giá sai,,,0,,,,,,\nBAD-02,,,,,abc,,,,,,\n"
    r = upload(client, owner_h, csv_text)
    assert r.status_code == 400
    detail = r.json()["detail"]
    assert [e["line"] for e in detail["errors"]] == [3, 4]
    assert client.get("/api/products", params={"q": "OK-01"}, headers=owner_h).json()["total"] == 0


def test_import_duplicates_inside_file(client, owner_h):
    csv_text = TEMPLATE_HEADER + "DUP-1,893000000001,A,,,100000,,,,,,\nDUP-1,,B,,,100000,,,,,,\nDUP-2,893000000001,C,,,100000,,,,,,\n"
    errors = upload(client, owner_h, csv_text, dry_run=True).json()["errors"]
    assert {e["line"] for e in errors} == {3, 4}


def test_import_semicolon_and_vietnamese_headers(client, owner_h):
    csv_text = "Mã;Tên;Giá bán;Nhóm hàng\nSC-77;Sạc 65W GaN;690000;Sạc\n"
    assert upload(client, owner_h, csv_text).json()["created"] == 1


def test_import_rejects_binary_and_staff(client, owner_h, staff_h):
    r = client.post("/api/products/import", headers=owner_h, files={"file": ("x.csv", b"\x89PNG\r\n\x1a\n\x00\x00", "text/csv")})
    assert r.status_code == 400
    assert upload(client, staff_h, TEMPLATE_HEADER).status_code == 403


def test_import_missing_columns(client, owner_h):
    r = upload(client, owner_h, "ten,gia_ban\nA,1000\n")
    assert r.status_code == 400 and "sku" in r.json()["detail"]


def xlsx(rows) -> bytes:
    from io import BytesIO

    from openpyxl import Workbook
    wb = Workbook()
    for r in rows:
        wb.active.append(r)
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_import_from_excel(client, owner_h):
    data = xlsx([["sku", "ten", "nhom_hang", "gia_ban", "gia_von", "theo_serial", "ma_vach"],
                 ["DT-XM-14", "Xiaomi 14", "Điện thoại", 15990000, 12000000.0, 1, "8938505970099"],
                 [None, None, None, None, None, None, None],  # dòng trống bị bỏ qua
                 ["PK001", "Trùng SKU", None, 100000, None, 0, None]])
    files = {"file": ("san pham.xlsx", data, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    preview = client.post("/api/products/import", params={"dry_run": True}, headers=owner_h, files=files).json()
    assert preview["to_create"] == 1 and preview["skipped"][0]["sku"] == "PK001"
    r = client.post("/api/products/import", headers=owner_h, files=files)
    assert r.status_code == 200 and r.json()["created"] == 1
    p = client.get("/api/products", params={"q": "DT-XM-14"}, headers=owner_h).json()["items"][0]
    assert p["sale_price"] == 15_990_000 and p["cost_price"] == 12_000_000 and p["track_serial"]


def test_excel_template_can_be_imported_as_is(client, owner_h):
    r = client.get("/api/products/import-template", params={"format": "xlsx"}, headers=owner_h)
    assert r.status_code == 200 and r.content[:2] == b"PK"
    files = {"file": ("mau.xlsx", r.content, "application/octet-stream")}
    body = client.post("/api/products/import", params={"dry_run": True}, headers=owner_h, files=files).json()
    assert body["errors"] == [] and body["preview"][0]["sku"] == "TN-SONY-02" and body["preview"][0]["sale_price"] == 8_490_000


def test_import_image_links(client, owner_h):
    from openpyxl import Workbook
    from io import BytesIO
    wb = Workbook()
    ws = wb.active
    ws.append(["sku", "ten", "gia_ban", "link_anh"])
    ws.append(["IMG-1", "Có link", 100000, "https://cdn.example.com/a.jpg"])
    ws.append(["IMG-2", "Drive", 100000, "https://drive.google.com/file/d/1AbCdEfGhIjKlMnOp/view?usp=sharing"])
    ws.append(["IMG-3", "Siêu liên kết", 100000, "Xem ảnh"])
    ws["D4"].hyperlink = "https://cdn.example.com/c.png"  # chữ hiển thị khác địa chỉ: lấy địa chỉ
    ws.append(["IMG-4", "Không ảnh", 100000, None])
    buf = BytesIO()
    wb.save(buf)
    files = {"file": ("anh.xlsx", buf.getvalue(), "application/octet-stream")}
    assert client.post("/api/products/import", headers=owner_h, files=files).json()["created"] == 4
    urls = {p["code"]: p["image_url"] for p in client.get("/api/products", params={"q": "IMG-"}, headers=owner_h).json()["items"]}
    assert urls == {"IMG-1": "https://cdn.example.com/a.jpg",
                    "IMG-2": "https://drive.google.com/thumbnail?id=1AbCdEfGhIjKlMnOp&sz=w800",
                    "IMG-3": "https://cdn.example.com/c.png", "IMG-4": None}
    bad = upload(client, owner_h, "sku,ten,gia_ban,link_anh\nIMG-9,Sai,100000,anh.jpg\n", dry_run=True).json()
    assert bad["errors"][0]["line"] == 2 and "http" in bad["errors"][0]["error"]


def test_import_rejects_broken_excel_and_xls(client, owner_h):
    r = client.post("/api/products/import", headers=owner_h, files={"file": ("x.xlsx", b"PK\x03\x04hong", "application/zip")})
    assert r.status_code == 400 and ".xlsx" in r.json()["detail"]
    old = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 100
    r = client.post("/api/products/import", headers=owner_h, files={"file": ("x.xls", old, "application/vnd.ms-excel")})
    assert r.status_code == 400 and ".xls" in r.json()["detail"]


# ---------------- /api/inventory/* ----------------
# TC-STK-04 (SRS 11.3)
def test_inventory_adjust_and_movements(client, owner_h, staff_h):
    pid = product_id(client, owner_h, "PK001")
    r = client.post("/api/inventory/adjust", json={"product_id": pid, "new_stock": 9, "note": "Kiểm kê cuối tháng"}, headers=owner_h)
    assert r.status_code == 200 and r.json()["stock"] == 9
    mv = client.get("/api/inventory/movements", params={"product_id": pid}, headers=owner_h).json()["items"][0]
    assert mv["type"] == "adjust" and mv["change"] == -3
    assert client.post("/api/inventory/adjust", json={"product_id": pid, "new_stock": 1, "note": "x"}, headers=staff_h).status_code == 403
    assert client.get("/api/inventory/movements", headers=staff_h).status_code == 403


def test_inventory_low_stock_for_staff(client, staff_h):
    rows = client.get("/api/inventory/low-stock", headers=staff_h).json()
    assert any(r["code"] == "PK002" for r in rows)
    assert all("cost_price" not in r for r in rows)  # FR-PRD-08


# ---------------- /api/export/* ----------------
# TC-EXP-01 (SRS 11.3)
def test_export_paths(client, owner_h, staff_h):
    r = client.get("/api/export/revenue", params={"format": "csv"}, headers=owner_h)
    assert r.status_code == 200 and r.headers["content-type"].startswith("text/csv")
    assert client.get("/api/export/invoices", params={"format": "xlsx"}, headers=staff_h).status_code == 200
    assert client.get("/api/export/revenue", headers=staff_h).status_code == 403
