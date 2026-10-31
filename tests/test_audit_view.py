"""Trang nhật ký thao tác: tên dễ đọc của bản ghi bị tác động, thay đổi trước -> sau, tìm kiếm, bộ lọc, xuất Excel."""
from io import BytesIO

from openpyxl import load_workbook

from app.services.audit_view import changes


def change_price(client, h, code, price):
    pid = client.get("/api/products", params={"q": code}, headers=h).json()["items"][0]["id"]
    assert client.put(f"/api/products/{pid}", json={"sale_price": price}, headers=h).status_code == 200


def test_changes_only_lists_fields_that_differ():
    assert changes('{"sale_price": 350000, "cost_price": 220000}', '{"sale_price": 390000, "cost_price": 220000}') == [
        {"field": "sale_price", "label": "Giá bán", "old": 350000, "new": 390000}]
    assert changes('{"role": "staff", "is_active": true}', '{"role": "owner", "is_active": false}') == [
        {"field": "role", "label": "Vai trò", "old": "Nhân viên bán hàng", "new": "Chủ cửa hàng"},
        {"field": "is_active", "label": "Trạng thái", "old": "Hoạt động", "new": "Đã khóa"}]
    assert changes(None, '{"points": 50}') == [{"field": "points", "label": "Điểm", "old": None, "new": 50}]
    assert changes("SN-1", "SN-2") == [{"field": None, "label": None, "old": "SN-1", "new": "SN-2"}]
    assert changes(None, None) == []


def test_items_have_readable_target_labels_and_changes(client, owner_h, admin_h):
    change_price(client, owner_h, "PK001", 390_000)
    client.post("/api/auth/login", json={"username": "khong-ton-tai", "password": "x"})
    body = client.get("/api/audit-logs", headers=admin_h).json()
    price = next(i for i in body["items"] if i["action"] == "PRICE_CHANGE")
    assert price["target"] == "PK001 · Tai nghe Bluetooth A1" and price["entity_label"] == "Sản phẩm"
    assert price["action_label"] == "Đổi giá" and price["user_name"] == "Chủ"
    assert price["changes"] == [{"field": "sale_price", "label": "Giá bán", "old": 350000, "new": 390000}]
    failed = next(i for i in body["items"] if i["action"] == "LOGIN_FAILED")
    assert failed["target"] == "khong-ton-tai" and failed["user_name"] is None
    assert {"value": "PRICE_CHANGE", "label": "Đổi giá"} in body["actions"]
    assert {"value": "products", "label": "Sản phẩm"} in body["entities"]
    assert any(u["name"] == "Chủ" for u in body["users"])


def test_search_by_target_name_and_actor(client, owner_h, admin_h):
    change_price(client, owner_h, "PK001", 390_000)
    change_price(client, owner_h, "PK003", 200_000)
    hits = client.get("/api/audit-logs", params={"q": "Sạc nhanh"}, headers=admin_h).json()["items"]
    assert [i["target"] for i in hits] == ["PK003 · Sạc nhanh 20W"]
    assert client.get("/api/audit-logs", params={"q": "Chủ", "action": "PRICE_CHANGE"}, headers=admin_h).json()["total"] == 2


def test_branch_tab_counts_follow_other_filters(client, owner_h, admin_h):
    bid = client.post("/api/branches", json={"name": "Cơ sở A"}, headers=admin_h).json()["id"]
    owner_id = next(u["id"] for u in client.get("/api/users", headers=admin_h).json() if u["username"] == "owner")
    client.put(f"/api/users/{owner_id}", json={"branch_id": bid}, headers=admin_h)
    change_price(client, owner_h, "PK001", 390_000)
    body = client.get("/api/audit-logs", params={"action": "PRICE_CHANGE", "branch_id": 0}, headers=admin_h).json()
    assert body["total"] == 0 and body["branch_counts"] == {str(bid): 1}  # tab vẫn đếm mọi chi nhánh


def test_export_excel_uses_current_filters(client, owner_h, admin_h, staff_h):
    change_price(client, owner_h, "PK001", 390_000)
    r = client.get("/api/audit-logs/export", params={"action": "PRICE_CHANGE"}, headers=admin_h)
    assert r.status_code == 200 and r.headers["content-type"].startswith("application/vnd.openxmlformats")
    rows = list(load_workbook(BytesIO(r.content)).active.iter_rows(values_only=True))
    assert rows[0][:6] == ("Thời điểm", "Chi nhánh", "Người thực hiện", "Đối tượng", "Thao tác", "Trên")
    assert len(rows) == 2 and rows[1][4] == "Đổi giá" and rows[1][6] == "Giá bán: 350.000" and rows[1][7] == "Giá bán: 390.000"
    assert client.get("/api/audit-logs/export", headers=staff_h).status_code == 403
