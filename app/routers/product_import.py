"""Nhập danh mục sản phẩm từ tệp Excel (.xlsx) hoặc CSV (SRS 7.7, bảng 8.5 POST /api/products/import, NFR-SEC-10).

Quy trình: tải mẫu -> điền -> tải lên với dry_run=true để xem trước -> tải lên lần nữa để nhập thật.
Nhập theo kiểu "tất cả hoặc không": chỉ cần một dòng lỗi là không dòng nào được ghi, tránh nhập dở dang.
SKU đã có trong hệ thống được bỏ qua (sửa giá, thông tin ở màn hình Sản phẩm để có audit log đổi giá).
Tồn đầu kỳ không nhập ở đây mà bằng một phiếu nhập có nhà cung cấp "Tồn đầu kỳ" (SRS 7.7).
"""
import csv
import io
import re
import zipfile
from datetime import date, datetime

from fastapi import APIRouter, Depends, File, HTTPException, Query, UploadFile
from fastapi.responses import Response
from pydantic import ValidationError
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from app.database import get_db
from app.models import Category, Product, User
from app.routers.catalog import _apply_category_defaults
from app.schemas import ProductIn
from app.security import MANAGERS
from app.services import audit
from app.services.export import MEDIA, to_csv

router = APIRouter(prefix="/api", tags=["catalog"])

MAX_BYTES = 5 * 1024 * 1024  # NFR-SEC-10
MAX_ROWS = 5000
# Cột trong mẫu: (tên cột, tên trường, bắt buộc, ví dụ)
COLUMNS = [
    ("sku", "code", True, "TN-SONY-02"),
    ("ma_vach", "barcode", False, "8938505970012"),
    ("ten", "name", True, "Tai nghe Sony WH-1000XM5"),
    ("nhom_hang", "category_name", False, "Tai nghe"),
    ("hang", "brand", False, "Sony"),
    ("gia_ban", "sale_price", True, "8490000"),
    ("gia_von", "cost_price", False, "7200000"),
    ("vat", "vat_rate", False, "10"),
    ("bao_hanh_thang", "warranty_months", False, "12"),
    ("ton_toi_thieu", "min_stock", False, "3"),
    ("theo_serial", "track_serial", False, "1"),
    ("mo_ta", "description", False, "Chống ồn chủ động, pin 30 giờ"),
]
# Chấp nhận cả tên cột tiếng Việt có dấu / tiếng Anh thường gặp
ALIASES = {
    "sku": "code", "ma": "code", "ma_sp": "code", "mã": "code", "code": "code",
    "ma_vach": "barcode", "mã vạch": "barcode", "barcode": "barcode",
    "ten": "name", "tên": "name", "ten_san_pham": "name", "name": "name",
    "nhom_hang": "category_name", "nhóm": "category_name", "nhóm hàng": "category_name", "category": "category_name",
    "hang": "brand", "hãng": "brand", "brand": "brand",
    "gia_ban": "sale_price", "giá bán": "sale_price", "sale_price": "sale_price",
    "gia_von": "cost_price", "giá vốn": "cost_price", "gia_nhap": "cost_price", "cost_price": "cost_price",
    "vat": "vat_rate", "vat_rate": "vat_rate",
    "bao_hanh_thang": "warranty_months", "bảo hành": "warranty_months", "warranty_months": "warranty_months",
    "ton_toi_thieu": "min_stock", "tồn tối thiểu": "min_stock", "min_stock": "min_stock",
    "theo_serial": "track_serial", "serial": "track_serial", "track_serial": "track_serial",
    "mo_ta": "description", "mô tả": "description", "description": "description",
}
INT_FIELDS = {"sale_price", "cost_price", "vat_rate", "warranty_months", "min_stock"}


# Giải thích từng cột, in ở trang "Hướng dẫn" của tệp mẫu Excel
COLUMN_HELP = {
    "sku": "Mã sản phẩm, không trùng. SKU đã có trong hệ thống sẽ được bỏ qua",
    "ma_vach": "Mã vạch EAN/UPC (nếu có), không trùng",
    "ten": "Tên sản phẩm",
    "nhom_hang": "Tên nhóm hàng. Nhóm chưa có sẽ được tạo mới",
    "hang": "Thương hiệu",
    "gia_ban": "Giá bán (đồng), lớn hơn 0",
    "gia_von": "Giá vốn (đồng)",
    "vat": "Thuế suất VAT %. Bỏ trống: lấy theo nhóm hàng",
    "bao_hanh_thang": "Số tháng bảo hành. Bỏ trống: lấy theo nhóm hàng",
    "ton_toi_thieu": "Tồn tối thiểu để cảnh báo sắp hết",
    "theo_serial": "1 = quản lý theo số serial / IMEI, 0 = không",
    "mo_ta": "Mô tả ngắn",
}


def _xlsx_template() -> bytes:
    """Tệp mẫu Excel: trang "San pham" (dòng tiêu đề + 1 dòng ví dụ) và trang "Huong dan" giải thích các cột."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "San pham"
    ws.append([c[0] for c in COLUMNS])
    ws.append([int(c[3]) if c[1] in INT_FIELDS or c[1] == "track_serial" else c[3] for c in COLUMNS])
    for cell, col in zip(ws[1], COLUMNS):
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="D4145A" if col[2] else "2563EB")  # cột bắt buộc tô hồng
        ws.column_dimensions[cell.column_letter].width = max(14, len(str(col[3])) + 4)
    for cell in ws[2]:
        if isinstance(cell.value, int) and cell.value >= 1000:
            cell.number_format = "#,##0"
    ws.cell(row=2, column=2).number_format = "@"  # mã vạch giữ dạng chữ, không thành số mũ
    ws.freeze_panes = "A2"
    guide = wb.create_sheet("Huong dan")
    guide.append(["Cột", "Bắt buộc", "Ý nghĩa"])
    for cell in guide[1]:
        cell.font = Font(bold=True)
    for name, _, required, _ in COLUMNS:
        guide.append([name, "Có" if required else "", COLUMN_HELP[name]])
    guide.append([])
    guide.append(["Lưu ý", "", "Giữ nguyên dòng tiêu đề. Chỉ trang đầu tiên được đọc. Tồn đầu kỳ nhập bằng phiếu nhập kho."])
    guide.column_dimensions["A"].width = 18
    guide.column_dimensions["C"].width = 80
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@router.get("/products/import-template")
def import_template(format: str = Query("csv", pattern="^(csv|xlsx)$"), _: User = Depends(MANAGERS)):
    if format == "xlsx":
        return Response(_xlsx_template(), media_type=MEDIA["xlsx"],
                        headers={"Content-Disposition": 'attachment; filename="mau_nhap_san_pham.xlsx"'})
    content = to_csv([c[0] for c in COLUMNS], [[c[3] for c in COLUMNS]])
    return Response(content, media_type="text/csv; charset=utf-8",
                    headers={"Content-Disposition": 'attachment; filename="mau_nhap_san_pham.csv"'})


def _cell(v) -> str:
    """Giá trị ô Excel thành chuỗi như khi gõ: 8490000.0 -> '8490000', ngày -> 'dd/mm/yyyy'."""
    if v is None:
        return ""
    if isinstance(v, bool):
        return "1" if v else "0"
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    if isinstance(v, (datetime, date)):
        return v.strftime("%d/%m/%Y")
    return str(v).strip()


def _xlsx_rows(data: bytes) -> list[list[str]]:
    """Đọc trang đầu tiên của tệp .xlsx (giá trị đã tính của công thức)."""
    from openpyxl import load_workbook

    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except (zipfile.BadZipFile, KeyError, ValueError, OSError):
        raise HTTPException(400, "Không đọc được tệp Excel. Hãy lưu lại dưới dạng .xlsx (Excel Workbook)")
    try:
        rows = []
        for r in wb.worksheets[0].iter_rows(values_only=True):
            if len(rows) > MAX_ROWS + 1:
                break  # quá giới hạn: parse_table báo lỗi, không đọc tiếp
            rows.append([_cell(v) for v in r])
        return rows
    finally:
        wb.close()


def _table(data: bytes, filename: str | None) -> list[list[str]]:
    """Bảng dữ liệu từ tệp tải lên. Loại tệp xác định bằng nội dung (NFR-SEC-10), không tin đuôi tên tệp."""
    if data[:4] == b"PK\x03\x04":  # .xlsx là tệp zip
        return _xlsx_rows(data)
    if data[:8] == b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1":  # Excel 97-2003 (.xls)
        raise HTTPException(400, "Tệp .xls (Excel 97-2003) chưa được hỗ trợ. Hãy mở bằng Excel và lưu lại dạng .xlsx")
    return _csv_rows(_decode(data))


def _decode(data: bytes) -> str:
    if b"\x00" in data[:4096]:
        raise HTTPException(400, "Tệp không phải Excel (.xlsx) hoặc CSV")  # kiểm tra loại tệp bằng nội dung
    for enc in ("utf-8-sig", "cp1258", "cp1252"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    raise HTTPException(400, "Không đọc được tệp. Hãy lưu CSV với mã hóa UTF-8")


def _number(raw: str) -> int | None:
    """'8.490.000', '8,490,000', '8490000 đ' -> 8490000; '1,5' -> 2 (làm tròn)."""
    s = re.sub(r"\s|₫|đ|vnd", "", raw.strip().lower())
    if not s:
        return None
    if re.fullmatch(r"\d{1,3}([.,]\d{3})+", s):  # dấu phân cách hàng nghìn
        s = re.sub(r"[.,]", "", s)
    return int(round(float(s.replace(",", "."))))


def _csv_rows(text: str) -> list[list[str]]:
    sample = text[:2048]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    return list(csv.reader(io.StringIO(text), dialect))


def parse_rows(text: str) -> tuple[list[dict], list[dict]]:
    return parse_table(_csv_rows(text))


def parse_table(table: list[list[str]]) -> tuple[list[dict], list[dict]]:
    """Dòng đầu là tiêu đề cột. Trả về (dòng hợp lệ, dòng lỗi); số dòng tính như trong tệp (tiêu đề là dòng 1)."""
    rows = [r for r in table if any(c.strip() for c in r)]
    if not rows:
        raise HTTPException(400, "Tệp rỗng")
    header = [ALIASES.get(h.strip().lower()) for h in rows[0]]
    if "code" not in header or "name" not in header or "sale_price" not in header:
        raise HTTPException(400, "Thiếu cột bắt buộc: sku, ten, gia_ban (tải tệp mẫu để xem đúng tên cột)")
    if len(rows) - 1 > MAX_ROWS:
        raise HTTPException(400, f"Tối đa {MAX_ROWS} dòng mỗi lần nhập")
    good, errors = [], []
    for line_no, r in enumerate(rows[1:], start=2):
        raw = {f: (r[i].strip() if i < len(r) else "") for i, f in enumerate(header) if f}
        try:
            fields: dict = {}
            for f, v in raw.items():
                if f in INT_FIELDS:
                    n = _number(v)
                    if n is not None:
                        fields[f] = n
                elif f == "track_serial":
                    fields[f] = v.lower() in ("1", "x", "co", "có", "yes", "true", "y")
                elif v:
                    fields[f] = v
            data = ProductIn(**fields)
        except (ValueError, ValidationError) as e:
            msg = "; ".join(str(x["msg"]).replace("Value error, ", "") for x in e.errors()) if isinstance(e, ValidationError) else "Số không hợp lệ"
            errors.append({"line": line_no, "sku": raw.get("code", ""), "error": msg})
            continue
        good.append({"line": line_no, "data": data})
    return good, errors


@router.post("/products/import")
async def import_products(file: UploadFile = File(...), dry_run: bool = Query(False),
                          db: Session = Depends(get_db), user: User = Depends(MANAGERS)):
    data = await file.read(MAX_BYTES + 1)
    if len(data) > MAX_BYTES:
        raise HTTPException(400, "Tệp quá 5 MB")
    good, errors = parse_table(_table(data, file.filename))

    # Trùng SKU / mã vạch trong tệp và với dữ liệu đã có
    existing_codes = {c.upper() for c in db.scalars(select(Product.code))}
    existing_barcodes = {b for b in db.scalars(select(Product.barcode).where(Product.barcode.is_not(None)))}
    seen_codes, seen_barcodes, to_create, skipped = set(), set(), [], []
    for row in good:
        d = row["data"]
        if d.code.upper() in existing_codes:
            skipped.append({"line": row["line"], "sku": d.code, "reason": "SKU đã có, bỏ qua"})
        elif d.code.upper() in seen_codes:
            errors.append({"line": row["line"], "sku": d.code, "error": "SKU bị lặp trong tệp"})
        elif d.barcode and (d.barcode in existing_barcodes or d.barcode in seen_barcodes):
            errors.append({"line": row["line"], "sku": d.code, "error": f"Mã vạch {d.barcode} đã được dùng"})
        else:
            seen_codes.add(d.code.upper())
            if d.barcode:
                seen_barcodes.add(d.barcode)
            to_create.append(row)
    errors.sort(key=lambda e: e["line"])
    preview = [{"line": r["line"], "sku": r["data"].code, "name": r["data"].name, "category": r["data"].category_name,
                "sale_price": r["data"].sale_price, "track_serial": r["data"].track_serial} for r in to_create[:200]]
    result = {"dry_run": dry_run, "to_create": len(to_create), "skipped": skipped, "errors": errors, "preview": preview}
    if dry_run or errors:
        if errors and not dry_run:
            raise HTTPException(400, {"message": f"Tệp có {len(errors)} dòng lỗi, chưa nhập dòng nào", **result})
        return result

    categories = {c.name.strip().lower(): c for c in db.scalars(select(Category))}
    for row in to_create:
        d = row["data"]
        fields = d.model_dump(exclude={"stock", "category_name", "category_id"})
        name = (d.category_name or "").strip()
        if name:
            cat = categories.get(name.lower())
            if cat is None:
                cat = Category(name=name)
                db.add(cat)
                db.flush()
                categories[name.lower()] = cat
            fields["category_id"] = cat.id
        _apply_category_defaults(db, fields)
        db.add(Product(**fields, stock=0))
    db.flush()
    audit.log(db, user, "PRODUCT_IMPORT", "products", None,
              new={"file": file.filename, "created": len(to_create), "skipped": len(skipped)})
    db.commit()
    total = db.scalar(select(func.count(Product.id)))
    return {**result, "created": len(to_create), "product_count": total}
